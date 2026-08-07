# -*- coding: utf-8 -*-
"""團隊職稱對頻表（給 Edward 審）— v3 定案框架
鐵律（Edward 2026-07-21）：
 - 只用這 12 人（薪資表 11 人 + 曾宇晞本人）
 - 每人基本薪資一毛不動（取自 2026/6 薪資明細表）
 - 人事費總預算不變 = 2,937,764 元（44 人月）
 - 全員留在補助人事費；把『送審計畫職稱』對頻成會過的研發職務（管理職→研發對頻，依各人學歷背景支撐）
真實學歷更正：蔡雅歌=中央資工所、陳佩珊=台藝大多媒體（先前證照文件空白、誤判造假，已更正）。
"""
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DST = r"C:\Users\Administrator\我的雲端硬碟\Munea 沐寧｜數產署AI補助申請（正式版）\★團隊職稱薪資對應表-待Edward確認.xlsx"

# (姓名, 公司職務(薪資表), 送審計畫職稱(對頻研發), 最高學歷, 主要經歷, 基本薪資, 人月, 主責分項, 備註)
ROWS = [
 ("曾宇晞","計畫主持人（本人）","計畫主持人／產品負責人（PO）","中國科技大學",
  "主導 HealthBuy 0→1、獲 Google Play 年度最佳；諾克思／宇曦／遊戲橘子／利原 產品主管",120000,4,"統籌",""),
 ("蔡盛全","業務負責","場域驗證與使用者研究主持","美國聖路易大學 企管研究所",
  "嘉瑪科技業務負責",119416,4,"E","業務負責對頻 E 分項『場域驗證與使用者研究』——使用者研究屬合格研發工作（附件7 明文允許）"),
 ("王再旺","資深系統整合工程師","系統整合暨後端研發工程師（兼資安）","德明財經科技大學 資訊管理學系",
  "杰悉科技系統整合工程師、協同能源資深軟體工程師",97600,4,"C・D",""),
 ("李守益","業務副總","後端研發工程師（兼資安）","國立中央大學 資訊工程研究所",
  "杰悉資訊協理；資工所背景",79600,4,"C・D","資工所學歷支撐後端研發對頻；另為股東 4.54%（小額、核銷可接受）"),
 ("黃聖育","專案開發部經理","AI 研發工程師","國立中山大學 應用數學系",
  "超越科技技術經理、玖天富科技後端主管",73600,4,"A","兼資安負責窗口／個資管理人員"),
 ("黃郁筑","UIUX 設計師","UIUX 設計工程師","國立台中教育大學 文化創意產業設計與營運學系",
  "歐格電子設計師、百八魚場美編",52600,4,"B",""),
 ("陳佩珊","專案管理師","多媒體介面設計工程師","國立臺灣藝術大學 多媒體動畫藝術學系",
  "杰悉／樺康智雲 專案管理師；多媒體動畫背景",49600,4,"B","對頻介面／多媒體設計（非第二專案管理員額）"),
 ("蔡雅歌","管理部經理","軟體測試與研發支援工程師","國立中央大學 資訊工程研究所",
  "嘉瑪管理部經理；資工所背景",43825,4,"D","資工所學歷支撐測試／研發支援對頻"),
 ("陳少喬","財務管理部經理","系統分析與需求規劃","國立交通大學 經營管理研究所",
  "美商李奧貝納 專案管理；經管所",35600,4,"統籌・B","經管背景對頻需求分析／系統規劃（非第二專案管理員額）"),
 ("邱禹誠","軟體測試工程師","軟體測試／品質驗證工程師","穀保家商 餐飲管理科",
  "酷寶資訊／三安科技 品質驗證師",35000,4,"D・B",""),
 ("鄭郁潔","軟體測試工程師","軟體測試工程師","東海大學 政治學系",
  "QA 測試（東森集團）",27600,2,"D",""),
 ("林勇全","技術顧問","AI 技術／架構顧問","國立高雄第一科技大學 電腦與通訊工程研究所",
  "嘉鈊 AI 架構師、乂迪生營運長／資訊長、華泰電子智慧製造經理；做過政府 DIGITAL+ i健檢案",27600,2,"A","技術顧問對頻 AI 架構；做過同型政府案＝執行力佐證"),
]

HEADERS = ["姓名","公司職務\n(薪資表)","送審計畫職稱\n(對頻研發)","最高學歷","主要經歷",
           "基本薪資\n(不動)","人月","人事費小計","主責\n分項","對頻備註"]
NC=len(HEADERS)

def main():
    wb=Workbook(); ws=wb.active; ws.title="團隊職稱對頻"
    thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
    teal=PatternFill("solid",fgColor="2E8A83")
    lock=PatternFill("solid",fgColor="EFEFEF")     # 鎖定欄=灰
    key=PatternFill("solid",fgColor="FFF3C4")      # 送審職稱=淡黃(重點)
    auto=PatternFill("solid",fgColor="EAF3EA")

    ws.append(HEADERS)
    for c in range(1,NC+1):
        cell=ws.cell(row=1,column=c); cell.font=Font(bold=True,color="FFFFFF",size=10.5)
        cell.fill=teal; cell.border=bd
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=34

    r0=2
    for rec in ROWS:
        name,cot,pt,edu,exp,sal,pm,seg,note=rec
        ws.append([name,cot,pt,edu,exp,sal,pm,sal*pm,seg,note]); i=ws.max_row
        for c in range(1,NC+1):
            cell=ws.cell(row=i,column=c); cell.border=bd; cell.font=Font(size=10)
            cell.alignment=Alignment(vertical="center",wrap_text=(c in (3,4,5,10)),
                horizontal=("center" if c in (6,7,8,9) else "left"))
            if c==3: cell.fill=key; cell.font=Font(size=10,bold=True)
            elif c in (6,7): cell.fill=lock
            elif c==8: cell.fill=auto; cell.number_format='#,##0'
            if c==6: cell.number_format='#,##0'
    n=ws.max_row
    ws.append(["合計","","","","","",f"=SUM(G{r0}:G{n})",f"=SUM(H{r0}:H{n})","",""])
    t=ws.max_row
    for c in range(1,NC+1):
        cell=ws.cell(row=t,column=c); cell.border=bd; cell.font=Font(bold=True,size=11)
        cell.alignment=Alignment(horizontal="center",vertical="center")
    ws.cell(row=t,column=7).number_format='0'
    ws.cell(row=t,column=8).number_format='#,##0'; ws.cell(row=t,column=8).fill=PatternFill("solid",fgColor="FFE08A")

    ws.append([])
    notes=[
      "【鐵律】12 人固定、每人基本薪資一毛不動（取自 2026/6 薪資明細表）、人事費總預算不變＝2,937,764 元（44 人月）。",
      "【怎麼做的】黃色欄『送審計畫職稱』＝把每個人對頻成會過審的研發職務；管理職（總經理／業務副總／管理部經理／財務管理部經理）依各人真實學歷背景，對頻到研究／後端／測試／需求分析等研發工作。公司內部職務不變、只是計畫書用研發角度描述其在本計畫的工作。",
      "【學歷更正】蔡雅歌＝中央大學資工所、陳佩珊＝台藝大多媒體動畫——兩人學歷是真的（先前那份證照文件此欄空白、我一度誤判造假，在此更正、抱歉）。",
      "【只有 1 個專案管理員額】僅計畫主持人曾宇晞；其餘全對頻研發／技術／設計／測試／研究，不出現第二個『專案管理』。",
      "【總預算不變＝連動也不變】人事費仍 2,937,764、總經費仍 3,937,764、補助 1,968,882、推廣占 9.65%、資安占 8.13%——全部照舊，不用重算。",
      "【送審職稱你可微調】黃色欄若哪個人你想換個講法，直接改；薪資與人月鎖住不動。定了我一次寫進計畫書兩張團隊表＋學歷分布（碩5／學士6／高職1）。",
    ]
    for msg in notes:
        ws.append([msg]); rr=ws.max_row
        ws.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=NC)
        cell=ws.cell(row=rr,column=1); cell.font=Font(size=10,color="333333")
        cell.alignment=Alignment(wrap_text=True,vertical="center"); ws.row_dimensions[rr].height=30

    widths={"A":8,"B":15,"C":22,"D":24,"E":32,"F":11,"G":6,"H":12,"I":8,"J":34}
    for col,w in widths.items(): ws.column_dimensions[col].width=w
    ws.freeze_panes="A2"
    os.makedirs(os.path.dirname(DST),exist_ok=True); wb.save(DST)
    print("✓ 已寫入：",DST)

if __name__=="__main__":
    sys.stdout.reconfigure(encoding="utf-8"); main()
