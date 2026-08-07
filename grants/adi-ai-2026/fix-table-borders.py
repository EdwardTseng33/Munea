# -*- coding: utf-8 -*-
"""為 build/ 內所有 docx 的表格加上標準框線（轉檔工具預設無框線 → Edward 7/17 抓到）。
sync-to-drive.py 會在轉檔後自動呼叫本工具。"""
import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.stdout.reconfigure(encoding="utf-8")

def add_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    # 移除舊的框線設定再加新的
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")       # 0.5pt 細線
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "444444")
        borders.append(el)
    tblPr.append(borders)

def fix(path: Path):
    doc = Document(str(path))
    n = 0
    for t in doc.tables:
        add_borders(t); n += 1
        # 巢狀表格也處理
        for row in t.rows:
            for cell in row.cells:
                for nt in cell.tables:
                    add_borders(nt); n += 1
    doc.save(str(path))
    return n

if __name__ == "__main__":
    build = Path(r"E:\Claude\Munea\grants\adi-ai-2026\build")
    total = 0
    for f in sorted(build.glob("*.docx")):
        n = fix(f)
        total += n
        print(f"✓ {f.name}：{n} 個表格加框線")
    print(f"\n完成：共處理 {total} 個表格")
