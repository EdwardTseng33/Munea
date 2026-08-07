#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""更新雲端那份進度表——只動指定的格子，其他一律原封不動。

為什麼不是「重生一份再覆蓋」：2026-08-10 差點踩到。Edward 會自己改那份表
（那天下午就把八月的完成線由 8/31 延到 9/5、把日期整批往後調），
拿本機重生的版本覆蓋等於把他的修改抹掉。所以流程一律是：
先下載雲端最新 → 只改該改的格子 → 寫回同一份。

用法：
    python update_progress.py --dry-run          # 只看會改什麼，不上傳
    python update_progress.py                    # 真的寫回
    python update_progress.py --file 我的更新.json

更新檔長這樣（updates.json）：
    {
      "note": "這次改了什麼，一句話",
      "rows": {
        "C1-2": {"狀態": "進行中", "完成%": 0.9, "備註": "……"},
        "D1-1": {"完成%": 0.8}
      }
    }
"""

import argparse
import io
import json
import os
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
FILE_ID = "10kTFpiDiDjelAwlvg9_P5gx5J6pA8aAi"
PROJECT = "gen-lang-client-0229303523"
SECRET = "munea-drive-sync-key"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_HINT = "8月"


def drive_service():
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    try:
        raw = subprocess.run(
            ["gcloud", "secrets", "versions", "access", "latest",
             "--secret", SECRET, "--project", PROJECT],
            capture_output=True, check=True, shell=(os.name == "nt"),
        ).stdout
    except subprocess.CalledProcessError as exc:
        detail = (exc.stderr or b"").decode("utf-8", "replace").strip().splitlines()
        raise SystemExit("拿不到鑰匙：" + (detail[-1] if detail else "未知原因"))

    creds = service_account.Credentials.from_service_account_info(
        json.loads(raw.decode("utf-8")), scopes=["https://www.googleapis.com/auth/drive"]
    )
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def download(service, dest):
    from googleapiclient.http import MediaIoBaseDownload

    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, service.files().get_media(fileId=FILE_ID))
    done = False
    while not done:
        _, done = downloader.next_chunk()
    dest.write_bytes(buf.getvalue())
    return dest


def apply_updates(path, rows):
    """回傳實際改了哪些格子。找不到的編號會回報，不會安靜跳過。"""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    name = next((s for s in wb.sheetnames if SHEET_HINT in s), None)
    if not name:
        raise SystemExit("找不到八月那個分頁，先確認表的結構有沒有被改過")
    ws = wb[name]

    header = {}
    for c in range(1, ws.max_column + 1):
        label = ws.cell(row=1, column=c).value
        if label:
            header[str(label).strip()] = c
    if "編號" not in header:
        raise SystemExit("八月分頁第一列找不到「編號」欄，表的結構跟預期不一樣")

    index = {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(row=r, column=header["編號"]).value
        if key:
            index[str(key).strip()] = r

    changed, missing = [], []
    for row_id, fields in rows.items():
        r = index.get(str(row_id).strip())
        if not r:
            missing.append(row_id)
            continue
        for col_name, value in fields.items():
            c = header.get(str(col_name).strip())
            if not c:
                missing.append(f"{row_id}／欄位「{col_name}」")
                continue
            cell = ws.cell(row=r, column=c)
            before = cell.value
            if before == value:
                continue
            cell.value = value
            changed.append((row_id, col_name, before, value))

    if missing:
        raise SystemExit("這些找不到，先確認編號或欄位名：" + "、".join(map(str, missing)))

    wb.save(path)
    return changed


def upload(service, path):
    from googleapiclient.http import MediaFileUpload

    media = MediaFileUpload(str(path), mimetype=XLSX_MIME, resumable=False)
    return service.files().update(
        fileId=FILE_ID, media_body=media, fields="id,name,modifiedTime,webViewLink"
    ).execute()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", default=str(ROOT / "updates.json"))
    ap.add_argument("--dry-run", action="store_true", help="只印出會改什麼，不寫回雲端")
    args = ap.parse_args()

    payload = json.loads(Path(args.file).read_text(encoding="utf-8"))
    rows = payload.get("rows") or {}
    if not rows:
        raise SystemExit("更新檔裡沒有 rows，沒東西可改")

    service = drive_service()
    work = ROOT / "out" / "cloud-working.xlsx"
    work.parent.mkdir(exist_ok=True)

    before = service.files().get(fileId=FILE_ID, fields="name,modifiedTime").execute()
    download(service, work)
    print(f"已抓下雲端最新版（上次修改 {before.get('modifiedTime')}）")

    changed = apply_updates(work, rows)
    if not changed:
        print("內容跟雲端一模一樣，不需要動。")
        return 0

    print(f"\n要改 {len(changed)} 格：")
    for row_id, col_name, old, new in changed:
        old_s = "(空白)" if old in (None, "") else str(old)[:34]
        print(f"  {row_id} · {col_name}：{old_s}  →  {str(new)[:44]}")

    if args.dry_run:
        print("\n(--dry-run：沒有寫回雲端)")
        return 0

    after = upload(service, work)
    print("\n已寫回同一份檔案（連結不變）")
    print("  檔名：", after.get("name"))
    print("  更新時間：", after.get("modifiedTime"))
    print("  連結：", after.get("webViewLink"))
    if payload.get("note"):
        print("  這次：", payload["note"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
